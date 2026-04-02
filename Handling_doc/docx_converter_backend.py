# ============================================
# PYTHON BACKEND FOR DOCX EDITING
# Better preservation of tables and formatting
# ============================================

from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from bs4 import BeautifulSoup
import io
import base64
import re

app = Flask(__name__)
CORS(app)

def html_to_docx(html_content):
    """
    Convert HTML to DOCX preserving tables and formatting
    """
    # Parse HTML
    soup = BeautifulSoup(html_content, 'html.parser')
    
    # Create Word document
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    # Process elements
    for element in soup.find_all(['p', 'h1', 'h2', 'h3', 'table', 'ul', 'ol']):
        
        if element.name == 'table':
            # Handle tables
            rows = element.find_all('tr')
            if not rows:
                continue
                
            # Create table in document
            table = doc.add_table(rows=len(rows), cols=len(rows[0].find_all(['td', 'th'])))
            table.style = 'Light Grid Accent 1'
            
            for i, row in enumerate(rows):
                cells = row.find_all(['td', 'th'])
                for j, cell in enumerate(cells):
                    doc_cell = table.rows[i].cells[j]
                    
                    # Get cell text with formatting
                    cell_text = cell.get_text(strip=True)
                    para = doc_cell.paragraphs[0]
                    
                    # Check for strong/bold
                    if cell.find(['strong', 'b']) or cell.name == 'th':
                        run = para.add_run(cell_text)
                        run.font.bold = True
                    else:
                        para.text = cell_text
                    
                    # Header row styling
                    if cell.name == 'th':
                        from docx.oxml.ns import qn
                        from docx.oxml import OxmlElement
                        
                        shading_elm = OxmlElement('w:shd')
                        shading_elm.set(qn('w:fill'), 'D9D9D9')
                        doc_cell._element.get_or_add_tcPr().append(shading_elm)
        
        elif element.name in ['h1', 'h2', 'h3']:
            # Handle headings
            level_map = {'h1': 0, 'h2': 1, 'h3': 2}
            text = element.get_text(strip=True)
            
            para = doc.add_paragraph(text)
            para.style = f'Heading {level_map[element.name] + 1}'
        
        elif element.name == 'p':
            # Handle paragraphs with inline formatting
            para = doc.add_paragraph()
            
            for content in element.children:
                if content.name is None:
                    # Plain text
                    text = str(content).strip()
                    if text:
                        para.add_run(text)
                
                elif content.name in ['strong', 'b']:
                    # Bold text
                    run = para.add_run(content.get_text(strip=True))
                    run.font.bold = True
                
                elif content.name in ['em', 'i']:
                    # Italic text
                    run = para.add_run(content.get_text(strip=True))
                    run.font.italic = True
                
                elif content.name == 'u':
                    # Underlined text
                    run = para.add_run(content.get_text(strip=True))
                    run.font.underline = True
                
                elif content.name == 'a':
                    # Hyperlink
                    text = content.get_text(strip=True)
                    url = content.get('href', '')
                    run = para.add_run(text)
                    run.font.color.rgb = RGBColor(0, 0, 255)
                    run.font.underline = True
                
                elif content.name == 'img':
                    # Image (base64)
                    src = content.get('src', '')
                    if src.startswith('data:image'):
                        try:
                            # Extract base64 data
                            image_data = re.search(r'base64,(.*)', src).group(1)
                            image_bytes = base64.b64decode(image_data)
                            
                            # Add image to document
                            image_stream = io.BytesIO(image_bytes)
                            para.add_run().add_picture(image_stream, width=Inches(4))
                        except Exception as e:
                            print(f"Image error: {e}")
        
        elif element.name in ['ul', 'ol']:
            # Handle lists
            for li in element.find_all('li', recursive=False):
                text = li.get_text(strip=True)
                para = doc.add_paragraph(text, style='List Bullet' if element.name == 'ul' else 'List Number')
    
    return doc


@app.route('/api/convert-html-to-docx', methods=['POST'])
def convert_html_to_docx():
    """
    API endpoint to convert HTML to DOCX
    """
    try:
        data = request.get_json()
        html_content = data.get('html', '')
        
        if not html_content:
            return jsonify({'error': 'No HTML content provided'}), 400
        
        # Convert HTML to DOCX
        doc = html_to_docx(html_content)
        
        # Save to bytes buffer
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        
        # Return as downloadable file
        return send_file(
            buffer,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            as_attachment=True,
            download_name='document.docx'
        )
    
    except Exception as e:
        print(f"Conversion error: {e}")
        return jsonify({'error': str(e)}), 500


@app.route('/api/docx-to-html', methods=['POST'])
def docx_to_html():
    """
    Convert DOCX to HTML (for editing)
    """
    try:
        file = request.files.get('file')
        if not file:
            return jsonify({'error': 'No file provided'}), 400
        
        # Read DOCX
        doc = Document(file)
        
        # Convert to HTML (simple version)
        html_parts = []
        
        for element in doc.element.body:
            if element.tag.endswith('p'):
                # Paragraph
                para_text = ''.join(node.text for node in element.iter() if hasattr(node, 'text'))
                html_parts.append(f'<p>{para_text}</p>')
            
            elif element.tag.endswith('tbl'):
                # Table
                html_parts.append('<table>')
                for row in element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr'):
                    html_parts.append('<tr>')
                    for cell in row.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc'):
                        cell_text = ''.join(node.text for node in cell.iter() if hasattr(node, 'text'))
                        html_parts.append(f'<td>{cell_text}</td>')
                    html_parts.append('</tr>')
                html_parts.append('</table>')
        
        html_content = ''.join(html_parts)
        
        return jsonify({'html': html_content})
    
    except Exception as e:
        print(f"Conversion error: {e}")
        return jsonify({'error': str(e)}), 500


@app.route('/health', methods=['GET'])
def health():
    return jsonify({'status': 'healthy'})


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)
